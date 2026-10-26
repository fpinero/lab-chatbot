"""Resource-bounded document/image preparation in a disposable subprocess."""

import json
import logging
import socket
import sys
import warnings
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from PIL import Image, ImageOps
from pypdf import PdfReader


class Rejected(Exception):
    pass


def deny_network(*args: object, **kwargs: object) -> None:
    raise Rejected("network_disabled")


def extract(path: Path, output: Path, extension: str, limits: dict[str, int]) -> dict[str, object]:
    text = ""
    notes: list[str] = []
    result: dict[str, object] = {"width": None, "height": None}
    if extension in {".png", ".jpg", ".jpeg", ".webp"}:
        Image.MAX_IMAGE_PIXELS = limits["pixels"]
        warnings.simplefilter("error", Image.DecompressionBombWarning)
        with Image.open(path, formats=["PNG", "JPEG", "WEBP"]) as image:
            expected = {".png": "PNG", ".jpg": "JPEG", ".jpeg": "JPEG", ".webp": "WEBP"}[extension]
            if image.format != expected or image.width * image.height > limits["pixels"]:
                raise Rejected("image_limits")
            if getattr(image, "n_frames", 1) != 1:
                raise Rejected("animated_image")
            image.load()
            normalized = ImageOps.exif_transpose(image)
            normalized = normalized.convert("RGBA" if "A" in normalized.getbands() else "RGB")
            if max(normalized.size) > limits["edge"]:
                normalized.thumbnail((limits["edge"], limits["edge"]))
                notes.append(f"Image resized to fit {limits['edge']} pixels per edge.")
            normalized.save(output, format="WEBP", quality=95, exif=b"", icc_profile=None)
            result.update(width=normalized.width, height=normalized.height)
            notes.append("Image normalized to WebP; metadata removed.")
        return {**result, "text": "", "status": "ready", "warnings": notes, "image": True}
    if extension == ".pdf":
        reader = PdfReader(path, strict=True)
        if reader.is_encrypted:
            raise Rejected("encrypted_document")
        if len(reader.pages) > limits["pages"]:
            raise Rejected("page_limit")
        chunks = []
        length = 0
        for page in reader.pages:
            content = page.get_contents()
            if content and len(content.get_data()) > 8 * 1024 * 1024:
                raise Rejected("processing_limit")
            value = page.extract_text() or ""
            length += len(value) + 1
            if length > limits["characters"]:
                raise Rejected("character_limit")
            chunks.append(value)
        text = "\n".join(chunks)
    elif extension == ".docx":
        with ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 1000 or sum(item.file_size for item in entries) > 64 * 1024 * 1024:
                raise Rejected("archive_limit")
            names = {item.filename for item in entries}
            if (
                len(names) != len(entries)
                or "word/document.xml" not in names
                or "[Content_Types].xml" not in names
            ):
                raise Rejected("invalid_document")
            for item in entries:
                name = item.filename.replace("\\", "/")
                if name.startswith("/") or ".." in name.split("/") or item.flag_bits & 1:
                    raise Rejected("unsafe_archive")
                if (
                    item.file_size > 20 * 1024 * 1024
                    or item.file_size > max(1, item.compress_size) * 100
                ):
                    raise Rejected("archive_limit")
                if any(value in name.lower() for value in ["vbaproject", "activex", "embeddings/"]):
                    raise Rejected("unsafe_document")
                if name.endswith((".xml", ".rels")):
                    xml_content = archive.read(item)
                    if any(
                        marker in xml_content.upper()
                        for marker in [b"<!DOCTYPE", b"<!ENTITY", b"MACROENABLED"]
                    ):
                        raise Rejected("unsafe_document")
        document = Document(str(path))
        chunks = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            chunks.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
        text = "\n".join(chunks)
        notes.append("Text in drawings or other complex layouts may not be extracted.")
    else:
        text = path.read_text(encoding="utf-8-sig")
        if "\x00" in text or text.lstrip().lower().startswith(("<!doctype html", "<html", "<svg")):
            raise Rejected("invalid_text")
    if len(text) > limits["characters"]:
        raise Rejected("character_limit")
    status = "ready" if text.strip() else "empty"
    if status == "empty":
        notes.append(
            "No selectable text was found. OCR is unavailable; this file cannot supply context."
        )
    return {**result, "text": text, "status": status, "warnings": notes, "image": False}


def main() -> None:
    logging.disable(logging.CRITICAL)
    try:
        import resource

        resource.setrlimit(resource.RLIMIT_CPU, (6, 6))
        if sys.platform.startswith("linux"):
            resource.setrlimit(resource.RLIMIT_DATA, (512 * 1024 * 1024, 512 * 1024 * 1024))
            resource.setrlimit(resource.RLIMIT_AS, (512 * 1024 * 1024, 512 * 1024 * 1024))
    except ImportError:
        pass
    socket.socket = deny_network  # type: ignore[misc, assignment]
    try:
        value = extract(Path(sys.argv[1]), Path(sys.argv[2]), sys.argv[3], json.loads(sys.argv[4]))
        print(json.dumps({"ok": True, **value}))
    except Rejected as exc:
        print(json.dumps({"ok": False, "code": str(exc)}))
    except (Image.DecompressionBombError, Image.DecompressionBombWarning):
        print(json.dumps({"ok": False, "code": "image_limits"}))
    except Exception:
        print(json.dumps({"ok": False, "code": "invalid_document"}))


if __name__ == "__main__":
    main()
